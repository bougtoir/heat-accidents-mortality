#!/usr/bin/env python3
"""
Convert the base manuscript into a Journal of Safety Research (JSR)
submission package.

Re-uses make_manuscript.py so every reported number is read from the generated
result CSVs.  The post-processing steps enforce JSR formatting (Elsevier /
National Safety Council guide for authors):
  - injury-safety / road-user framing in title and abstract
  - structured abstract (Introduction, Method, Results, Conclusions,
    Practical Applications; <=300 words) + keywords
  - APA (author-date) in-text citations and an alphabetical reference list
    with all authors listed
  - "Practical Applications" section placed at the end of the main text
  - separate editable figure files + legends-only manuscript
  - JSR cover letter and Highlights file
"""
import os
import re
import shutil
import sys
import tempfile

from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
BW = os.environ.get("FIGURES_BW") == "1"
SUFFIX = "_bw" if BW else ""
DEFAULT_FIG = os.path.join(ROOT, "output", "figures_bw" if BW else "figures")
DEFAULT_MAN = os.path.join(ROOT, "output", "manuscript_bw" if BW else "manuscript")
os.environ.setdefault("FIGURES_DIR", DEFAULT_FIG)
os.environ.setdefault("MANUSCRIPT_DIR", DEFAULT_MAN)

sys.path.insert(0, HERE)
import make_manuscript as mm  # noqa: E402
MAN = mm.MAN

JSR_TITLE = ("Hotter-than-normal days and US traffic-crash deaths: "
             "road-user-specific injury risk and implications for heat-aware "
             "road safety")

JSR_KEYWORDS = ("extreme heat; traffic safety; road traffic injury; "
                "road-user type; climate adaptation; distributed-lag model")

# ---------------------------------------------------------------------------
# APA (author-date) in-text citations and alphabetical references.
# JSR requires APA style with all authors listed alphabetically.
# ---------------------------------------------------------------------------
JSR_INTEXT = {
    "dlnm": "Gasparrini, Armstrong, & Kenward, 2010",
    "lancet": "Gasparrini et al., 2015",
    "basu": "Basu, 2009",
    "fars": "National Highway Traffic Safety Administration [NHTSA], n.d.",
    "ghcn": "Menne et al., 2012",
    "cdc": "Centers for Disease Control and Prevention [CDC], n.d.",
    "npa": "National Police Agency of Japan [NPA], n.d.",
    "attrib": "Gasparrini & Leone, 2014",
    "eia": "US Energy Information Administration [EIA], n.d.",
    "fhwa": "Federal Highway Administration [FHWA], n.d.-a",
    "fhwa_vm2": "Federal Highway Administration [FHWA], n.d.-b",
    "census_pep": "US Census Bureau, n.d.",
    "ipcc": "IPCC, 2021",
    "daanen": "Daanen, van de Vliert, & Huang, 2003",
    "liang2022": "Liang et al., 2022",
    "liang2021_aap": "Liang et al., 2021",
    "care": "European Commission, n.d.",
}

# Alphabetical, all authors listed (APA).
JSR_REF_TEXT = {
    "basu": (
        "Basu, R. (2009). High ambient temperature and mortality: a review of "
        "epidemiologic studies from 2001 to 2008. Environmental Health, 8, 40. "
        "https://doi.org/10.1186/1476-069X-8-40"
    ),
    "cdc": (
        "Centers for Disease Control and Prevention. (n.d.). CDC WONDER Underlying "
        "Cause of Death database. https://wonder.cdc.gov/"
    ),
    "daanen": (
        "Daanen, H. A. M., van de Vliert, E., & Huang, X. (2003). Driving performance "
        "in cold, warm, and thermoneutral environments. Applied Ergonomics, 34(6), "
        "597-602. https://doi.org/10.1016/S0003-6870(03)00055-3"
    ),
    "care": (
        "European Commission, Directorate-General for Mobility and Transport. (n.d.). "
        "Community database on road accidents (CARE). "
        "https://road-safety.transport.ec.europa.eu/european-road-safety-observatory/"
        "methodology-and-research/care-database_en"
    ),
    "fhwa": (
        "Federal Highway Administration. (n.d.-a). Traffic Volume Trends. "
        "https://www.fhwa.dot.gov/policyinformation/travel_monitoring/tvt.cfm"
    ),
    "fhwa_vm2": (
        "Federal Highway Administration. (n.d.-b). Highway Statistics VM-2: Annual "
        "state vehicle-miles travelled. "
        "https://www.fhwa.dot.gov/policyinformation/statistics.cfm"
    ),
    "dlnm": (
        "Gasparrini, A., Armstrong, B., & Kenward, M. G. (2010). Distributed lag "
        "non-linear models. Statistics in Medicine, 29(21), 2224-2234. "
        "https://doi.org/10.1002/sim.3940"
    ),
    "attrib": (
        "Gasparrini, A., & Leone, M. (2014). Attributable risk from distributed lag "
        "models. BMC Medical Research Methodology, 14, 55. "
        "https://doi.org/10.1186/1471-2288-14-55"
    ),
    "lancet": (
        "Gasparrini, A., Guo, Y., Hashizume, M., Lavigne, E., Zanobetti, A., "
        "Schwartz, J., Tobias, A., Tong, S., Rocklov, J., Forsberg, B., Leone, M., "
        "De Sario, M., Bell, M. L., Guo, Y.-L., Wu, C.-F., Kan, H., Yi, S.-M., "
        "de Sousa Zanotti Stagliorio Coelho, M., Saldiva, P. H. N., Honda, Y., "
        "Kim, H., & Armstrong, B. (2015). Mortality risk attributable to high and "
        "low ambient temperature: a multicountry observational study. The Lancet, "
        "386(9991), 369-375. https://doi.org/10.1016/S0140-6736(14)62114-0"
    ),
    "ipcc": (
        "IPCC. (2021). Climate Change 2021: The Physical Science Basis. Contribution "
        "of Working Group I to the Sixth Assessment Report of the Intergovernmental "
        "Panel on Climate Change. Cambridge University Press."
    ),
    "liang2021_aap": (
        "Liang, M., Zhao, D., Wu, Y., Ye, P., Wang, Y., Yao, Z., Xu, Y., Zhang, Y., "
        "Zhou, Y., & Pan, X. (2021). Short-term effects of ambient temperature and "
        "road traffic accident injuries in Dalian, Northern China: A distributed lag "
        "non-linear analysis. Accident Analysis & Prevention, 153, 106057. "
        "https://doi.org/10.1016/j.aap.2021.106057"
    ),
    "liang2022": (
        "Liang, M., Min, M., Guo, X., Song, Q., Wang, H., Li, N., Zhao, D., Xu, Y., "
        "& Pan, X. (2022). The relationship between ambient temperatures and road "
        "traffic injuries: a systematic review and meta-analysis. Environmental "
        "Science and Pollution Research, 29(33), 50647-50660. "
        "https://doi.org/10.1007/s11356-022-19437-y"
    ),
    "ghcn": (
        "Menne, M. J., Durre, I., Vose, R. S., Gleason, B. E., & Houston, T. G. "
        "(2012). An overview of the Global Historical Climatology Network-Daily "
        "database. Journal of Atmospheric and Oceanic Technology, 29(7), 897-910. "
        "https://doi.org/10.1175/JTECH-D-11-00103.1"
    ),
    "fars": (
        "National Highway Traffic Safety Administration. (n.d.). Fatality Analysis "
        "Reporting System (FARS). "
        "https://www.nhtsa.gov/research-data/fatality-analysis-reporting-system-fars"
    ),
    "npa": (
        "National Police Agency of Japan. (n.d.). Traffic accident statistics open "
        "data. https://www.npa.go.jp/publications/statistics/koutsuu/opendata/"
    ),
    "eia": (
        "US Energy Information Administration. (n.d.). Weekly finished motor "
        "gasoline product supplied (PET.WGFUPUS2.W). https://www.eia.gov/"
    ),
    "census_pep": (
        "US Census Bureau. (n.d.). Population Estimates Program: Annual state "
        "population estimates. https://www.census.gov/programs-surveys/popest.html"
    ),
}

JSR_HIGHLIGHTS = [
    "Hotter-than-normal days carry an acute excess of US crash deaths.",
    "The excess concentrates in open-air road users, not vehicle occupants.",
    "Heat-attributable crash deaths rival recorded direct-heat deaths.",
    "Warming projections imply more US crash deaths per year per degree.",
    "Heat-aware road safety and shared-mobility adaptation are warranted.",
]

JSR_COVER = """22 August 2026

The Editor-in-Chief, Journal of Safety Research

Dear Editor-in-Chief,

We submit for your consideration our manuscript, "Hotter-than-normal days and US traffic-crash deaths: road-user-specific injury risk and implications for heat-aware road safety", as a Research Article.

This study speaks directly to the Journal's traffic-safety scope. Using only public data, this ecological, population-level study shows that days hotter than the local seasonal norm carry an acute, same-day excess of US traffic-crash deaths that survives adjustment for driving activity, precipitation and heat-stress controls in a VIF-screened full-controls model. Crucially for injury prevention, the excess is graded by direct heat exposure: it is small for enclosed vehicle occupants but several-fold larger for motorcyclists, pedestrians and cyclists, and largest in the hottest hours of the day. The net heat-attributable burden is comparable in magnitude to all officially recorded direct-heat deaths yet remains invisible to both heat-mortality and road-safety surveillance because such deaths continue to be coded as ordinary crashes. Scenario projections suggest this hidden injury burden would grow under continued warming.

The findings translate into concrete safety practice: targeted heat warnings for open-air road users and outdoor delivery riders during hot periods, heat-aware operation of shared e-scooter and e-bike fleets (whose batteries, motors and tyres are heat-stressed and whose docking and charging points often lack shade), and integration of temperature forecasts into crash-prevention messaging consistent with a Safe System approach to environmental risk.

The manuscript is original, is not under consideration elsewhere, and all authors approve submission. All data are public and the complete analysis pipeline is openly available and fully reproducible at https://github.com/bougtoir/heat-accidents-mortality (make all for data and figures, then make jsr for this submission package), with no hard-coded results. We declare no competing interests. We confirm that this manuscript has not been posted to a preprint server.

We look forward to your consideration.

Yours sincerely,


Tatsuki Onishi, on behalf of all authors

Data Science and AI Innovation Research Promotion Center

Shiga University of Medical Science

Seta Tsukinowa-cho, Otsu, Shiga 520-2192, Japan

E-mail: bougtoir@gmail.com
"""


# ---------------------------------------------------------------------------
# docx helpers
# ---------------------------------------------------------------------------
def _heading_text(p):
    return (p.text or "").strip()


def _set_heading_text(p, text):
    p.text = text


def _remove_paragraph(p):
    p._element.getparent().remove(p._element)


def _is_heading(p):
    return bool(p.style and p.style.name.startswith("Heading"))


def _new_paragraph_after(doc, anchor, text="", bold_label=None):
    p = doc.add_paragraph()
    if bold_label:
        r = p.add_run(bold_label)
        r.bold = True
    if text:
        p.add_run(text)
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_after = Pt(6)
    anchor._element.addnext(p._element)
    return p


# ---------------------------------------------------------------------------
# APA conversion: numbered superscript citations -> (author, year) in text
# ---------------------------------------------------------------------------
def _intext_for_group(nums_str, ref_order):
    parts = []
    for n in nums_str.split(","):
        n = n.strip()
        if not n.isdigit():
            continue
        key = ref_order[int(n) - 1]
        parts.append(JSR_INTEXT[key])
    if not parts:
        return None
    return "(" + "; ".join(parts) + ")"


def _convert_citations_apa(doc, ref_order):
    """Replace every superscript citation run '1,2' with an APA in-text group."""
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.superscript and r.text and re.fullmatch(r"\d+(,\d+)*", r.text):
                rep = _intext_for_group(r.text, ref_order)
                if rep is None:
                    continue
                r.text = rep
                r.font.superscript = False


# ---------------------------------------------------------------------------
# Front-matter post-processing
# ---------------------------------------------------------------------------
def _replace_title(doc):
    for p in doc.paragraphs[:6]:
        for r in p.runs:
            if "under-recognised risk factor for US traffic-crash" in (r.text or ""):
                r.text = JSR_TITLE
                return


def _make_jsr_abstract(doc, abstract_heading, body_ps):
    for p in body_ps:
        _remove_paragraph(p)

    abstract_heading.text = "Abstract"

    intro = ("Ambient heat is an established mortality risk, but its contribution "
             "to road traffic injury deaths is under-recognised because crash "
             "decedents rarely undergo heat-oriented post-mortem assessment.")
    method = ("We fitted quasi-Poisson distributed-lag models of local seasonal "
              "temperature anomalies to US state-day (FARS, 2016-2022) and Japan "
              "prefecture-day (NPA, 2019-2024) panels of traffic-crash deaths, "
              "with activity, precipitation and heat-stress sensitivity controls.")
    results = (f"Among {int(float(mm.US['total_deaths'])):,} US crash deaths, a +9\u00b0C "
               f"anomaly raised same-day mortality, RR {mm.rr(mm.US, 'sameday_RR_anom+9C')}; "
               f"the VIF-screened full-controls model gave RR {mm.ctrlrr(mm.CTRL)}. "
               "The excess was graded by heat exposure: motorcyclists "
               f"RR {mm.f(mm.USER['motorcyclist']['sameday_RR_+9C'])}, pedestrians "
               f"RR {mm.f(mm.USER['pedestrian']['sameday_RR_+9C'])}, cyclists "
               f"RR {mm.f(mm.USER['cyclist']['sameday_RR_+9C'])}, versus vehicle "
               f"occupants RR {mm.f(mm.USER['vehicle_occupant']['sameday_RR_+9C'])}. "
               f"Net heat-attributable deaths averaged "
               f"{float(mm.US['net_heat_attributable_per_year']):.0f} per year, similar to "
               f"{mm.CDC_MEAN:.0f} recorded direct-heat deaths; +1\u00b0C warming projected "
               f"{float(mm.PROJ_D[1]['extra_deaths_per_year']):.0f} additional deaths/year. "
               "The Japanese estimate was imprecise.")
    conclusions = ("Unusually hot days carry an acute, road-user-specific excess of US "
                   "crash deaths comparable to recorded direct-heat mortality and "
                   "invisible to routine surveillance.")
    practical = ("Heat warnings targeted at open-air road users, heat-aware operation "
                 "of shared e-scooter and e-bike fleets, shaded docking and charging "
                 "infrastructure, and temperature-forecast-linked crash-prevention "
                 "messaging could reduce this uncounted injury burden.")
    order = [
        ("Introduction", intro),
        ("Method", method),
        ("Results", results),
        ("Conclusions", conclusions),
        ("Practical Applications", practical),
    ]
    total = sum(len(t.split()) for _l, t in order if t)
    if total > 300:
        print(f"WARNING: JSR abstract is {total} words (>300)")

    prev = abstract_heading
    for label, text in order:
        if not text:
            continue
        p = doc.add_paragraph()
        r = p.add_run(f"{label}: ")
        r.bold = True
        p.add_run(text)
        p.paragraph_format.line_spacing = 2.0
        p.paragraph_format.space_after = Pt(6)
        prev._element.addnext(p._element)
        prev = p

    kw = doc.add_paragraph()
    r = kw.add_run("Keywords: ")
    r.bold = True
    kw.add_run(JSR_KEYWORDS)
    kw.paragraph_format.line_spacing = 2.0
    kw.paragraph_format.space_after = Pt(6)
    prev._element.addnext(kw._element)


def _remove_research_in_context(doc, heading, body_ps):
    _remove_paragraph(heading)
    for p in body_ps:
        _remove_paragraph(p)


def _move_practical_applications(doc):
    """Move 'Implications for road safety' to the end of the main text, retitled
    'Practical Applications' (JSR requires this section last)."""
    paras = doc.paragraphs
    imp_idx = None
    concl_idx = None
    contrib_idx = None
    for i, p in enumerate(paras):
        if not _is_heading(p):
            continue
        t = _heading_text(p)
        if t == "Implications for road safety":
            imp_idx = i
        elif t == "Contributors":
            contrib_idx = i
        elif t == "Conclusion":
            concl_idx = i
    if imp_idx is None or contrib_idx is None:
        return
    # collect the section's elements (heading + following body paragraphs)
    block = []
    for p in paras[imp_idx:]:
        block.append(p)
        if p is not paras[imp_idx] and _is_heading(p):
            block.pop()
            break
    _set_heading_text(paras[imp_idx], "Practical Applications")
    anchor = paras[contrib_idx - 1] if concl_idx is not None else paras[contrib_idx]
    for p in block:
        anchor._element.addnext(p._element)
        anchor = p


def _rename_headings(doc):
    renames = {
        "Methods": "Materials and Methods",
        "Contributors": "Author Contributions",
        "Declaration of competing interests": "Declaration of competing interest",
        "Declaration of generative AI use":
            "Declaration of generative AI and AI-assisted technologies in the writing process",
        "Data availability": "Data availability",
    }
    for p in doc.paragraphs:
        if _is_heading(p):
            t = _heading_text(p)
            if t in renames:
                _set_heading_text(p, renames[t])


def _fix_data_availability(doc):
    for p in doc.paragraphs:
        if "make ehp for the EHP submission package" in (p.text or ""):
            p.text = p.text.replace(
                "make ehp for the EHP submission package",
                "make jsr for the JSR submission package",
            )
        if "make aap for the manuscript and submission package" in (p.text or ""):
            p.text = p.text.replace(
                "make aap for the manuscript and submission package",
                "make jsr for the JSR submission package",
            )


def _rewrite_references_apa(doc):
    """Replace the numbered reference list with an alphabetical APA list."""
    ref_heading = None
    for p in doc.paragraphs:
        if _is_heading(p) and _heading_text(p) == "References":
            ref_heading = p
            break
    if ref_heading is None:
        return
    cited_keys = list(mm._ref_order)
    # remove old numbered entries (all paragraphs after References heading)
    seen_refs = False
    for p in doc.paragraphs:
        if p._p is ref_heading._p:
            seen_refs = True
            continue
        if seen_refs:
            _remove_paragraph(p)
    prev = ref_heading
    for key in sorted(cited_keys, key=lambda k: JSR_REF_TEXT[k].lower()):
        p = doc.add_paragraph(JSR_REF_TEXT[key])
        p.paragraph_format.space_after = Pt(4)
        prev._element.addnext(p._element)
        prev = p


def _postprocess(src_path, dst_path):
    doc = Document(src_path)

    _replace_title(doc)
    _convert_citations_apa(doc, mm._ref_order)

    abstract_heading = None
    abstract_body = []
    research_heading = None
    research_body = []
    collect = None
    for p in doc.paragraphs:
        if _is_heading(p):
            txt = _heading_text(p)
            if txt == "Summary":
                abstract_heading = p
                collect = "abstract"
                continue
            if txt == "Research in context":
                research_heading = p
                collect = "research"
                continue
            collect = None
            continue
        if collect == "abstract":
            abstract_body.append(p)
        elif collect == "research":
            research_body.append(p)

    if abstract_heading and abstract_body:
        _make_jsr_abstract(doc, abstract_heading, abstract_body)
    if research_heading:
        _remove_research_in_context(doc, research_heading, research_body)

    _move_practical_applications(doc)
    _rename_headings(doc)
    _fix_data_availability(doc)
    _rewrite_references_apa(doc)

    doc.save(dst_path)
    print("wrote", dst_path)


# ---------------------------------------------------------------------------
# Cover letter, highlights and submission bundle
# ---------------------------------------------------------------------------
def _build_cover_letter(path):
    doc = Document()
    mm.setup(doc)
    for block in JSR_COVER.strip().split("\n\n"):
        text = " ".join(block.splitlines())
        if text:
            p = doc.add_paragraph(text)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(6)
    doc.save(path)
    print("wrote", path)


def _build_highlights(path):
    doc = Document()
    mm.setup(doc)
    doc.add_heading("Highlights", 1)
    for b in JSR_HIGHLIGHTS:
        assert len(b) <= 85, f"highlight >85 chars: {b!r}"
        doc.add_paragraph(b, style="List Bullet")
    doc.save(path)
    print("wrote", path)


def _build_figures_stage():
    dest_dir = tempfile.mkdtemp(prefix="jsr_figures_", dir=MAN)
    src = os.path.join(MAN, "submission_figures")
    if not os.path.isdir(src):
        raise FileNotFoundError(f"{src} not found; run 'make manuscript' first.")
    shutil.copytree(src, os.path.join(dest_dir, f"jsr_submission_figures{SUFFIX}"))
    return dest_dir


def _build_submission_zip(fig_stage):
    zip_base = os.path.join(MAN, f"jsr_submission_package{SUFFIX}")
    stage = zip_base + "_stage"
    if os.path.exists(stage):
        shutil.rmtree(stage)
    os.makedirs(stage)
    files = [
        f"heat_crash_mortality_jsr{SUFFIX}.docx",
        f"heat_crash_mortality_jsr_legends{SUFFIX}.docx",
        f"jsr_cover_letter{SUFFIX}.docx",
        f"highlights{SUFFIX}.docx",
        f"tables{SUFFIX}.docx",
        f"figures{SUFFIX}.pptx",
        f"strobe_checklist{SUFFIX}.docx",
    ]
    for name in files:
        src = os.path.join(MAN, name)
        if not os.path.exists(src) and SUFFIX:
            src = os.path.join(MAN, name.replace(SUFFIX, "", 1))
        if os.path.exists(src):
            shutil.copyfile(src, os.path.join(stage, name))
    if fig_stage and os.path.isdir(fig_stage):
        for entry in os.listdir(fig_stage):
            entry_src = os.path.join(fig_stage, entry)
            entry_dst = os.path.join(stage, entry)
            if os.path.isdir(entry_src):
                shutil.copytree(entry_src, entry_dst)
            else:
                shutil.copyfile(entry_src, entry_dst)
    if os.path.exists(zip_base + ".zip"):
        os.remove(zip_base + ".zip")
    shutil.make_archive(zip_base, "zip", stage)
    shutil.rmtree(stage)
    print("wrote", zip_base + ".zip")


def main():
    os.makedirs(MAN, exist_ok=True)

    mm.build_pptx()
    mm.build_tables_docx()
    mm.build_strobe()
    mm.build_submission_figures()

    tmpdir = tempfile.mkdtemp(prefix="jsr_build_", dir=MAN)
    original_man = mm.MAN
    mm.MAN = tmpdir
    try:
        mm.build_manuscript("tmp_inline.docx", embed=True)
        mm.build_manuscript("tmp_legends.docx", embed=False)
    finally:
        mm.MAN = original_man

    inline_src = os.path.join(tmpdir, "tmp_inline.docx")
    legends_src = os.path.join(tmpdir, "tmp_legends.docx")

    # Capture the citation order from the last build (both builds share it).
    _postprocess(inline_src, os.path.join(MAN, f"heat_crash_mortality_jsr{SUFFIX}.docx"))
    _postprocess(legends_src, os.path.join(MAN, f"heat_crash_mortality_jsr_legends{SUFFIX}.docx"))

    shutil.rmtree(tmpdir)

    _build_cover_letter(os.path.join(MAN, f"jsr_cover_letter{SUFFIX}.docx"))
    _build_highlights(os.path.join(MAN, f"highlights{SUFFIX}.docx"))

    fig_stage = _build_figures_stage()
    try:
        _build_submission_zip(fig_stage)
    finally:
        shutil.rmtree(fig_stage)


if __name__ == "__main__":
    main()
